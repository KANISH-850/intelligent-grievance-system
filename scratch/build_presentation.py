import os
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def create_presentation():
    prs = Presentation()
    
    # 16:9 Widescreen dimensions
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    
    # Color Palette
    BG_COLOR = RGBColor(248, 250, 252)       # Slate 50 (#F8FAFC)
    CARD_BG = RGBColor(255, 255, 255)        # Pure White
    CARD_BORDER = RGBColor(226, 232, 240)    # Slate 200
    NAVY_PRIMARY = RGBColor(15, 23, 42)      # Slate 900 (#0F172A)
    BLUE_ACCENT = RGBColor(29, 78, 216)      # Blue 700 (#1D4ED8)
    TEXT_MAIN = RGBColor(30, 41, 59)         # Slate 800 (#1E293B)
    TEXT_MUTED = RGBColor(100, 116, 139)    # Slate 500 (#64748B)
    GREEN_SUCCESS = RGBColor(22, 101, 52)   # Green 800
    GOLD_AMBER = RGBColor(180, 83, 9)       # Amber 700
    
    blank_layout = prs.slide_layouts[6] # Blank slide
    
    def set_slide_background(slide):
        background = slide.background
        fill = background.fill
        fill.solid()
        fill.fore_color.rgb = BG_COLOR

    def add_header(slide, title_text, category_text="ACADEMIC PROJECT DEFENSE"):
        txBox = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.733), Inches(0.9))
        tf = txBox.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        
        p0 = tf.paragraphs[0]
        p0.text = category_text.upper()
        p0.font.size = Pt(10)
        p0.font.bold = True
        p0.font.color.rgb = BLUE_ACCENT
        p0.font.name = "Arial"
        p0.space_after = Pt(2)
        
        p1 = tf.add_paragraph()
        p1.text = title_text
        p1.font.size = Pt(22)
        p1.font.bold = True
        p1.font.color.rgb = NAVY_PRIMARY
        p1.font.name = "Arial"

    def add_card(slide, left, top, width, height, title=None, title_color=NAVY_PRIMARY, bg_color=CARD_BG, border_color=CARD_BORDER):
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height))
        shape.fill.solid()
        shape.fill.fore_color.rgb = bg_color
        if border_color:
            shape.line.color.rgb = border_color
            shape.line.width = Pt(1)
        else:
            shape.line.fill.background()
            
        if title:
            txBox = slide.shapes.add_textbox(Inches(left + 0.2), Inches(top + 0.15), Inches(width - 0.4), Inches(0.4))
            tf = txBox.text_frame
            tf.word_wrap = True
            tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
            p = tf.paragraphs[0]
            p.text = title
            p.font.size = Pt(14)
            p.font.bold = True
            p.font.color.rgb = title_color
            p.font.name = "Arial"
        return shape

    # ----------------------------------------------------
    # SLIDE 1: Title Slide
    # ----------------------------------------------------
    s1 = prs.slides.add_slide(blank_layout)
    set_slide_background(s1)
    
    add_card(s1, 0.8, 0.8, 11.733, 5.9, bg_color=CARD_BG, border_color=CARD_BORDER)
    
    tbox = s1.shapes.add_textbox(Inches(1.2), Inches(1.2), Inches(10.933), Inches(5.1))
    tf = tbox.text_frame
    tf.word_wrap = True
    
    p = tf.paragraphs[0]
    p.text = "FINAL ACADEMIC PROJECT DEFENSE"
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = BLUE_ACCENT
    p.space_after = Pt(14)
    
    p = tf.add_paragraph()
    p.text = "Intelligent Multilingual Grievance Categorization &\nAutomated Dispatch Framework"
    p.font.size = Pt(28)
    p.font.bold = True
    p.font.color.rgb = NAVY_PRIMARY
    p.space_after = Pt(10)
    
    p = tf.add_paragraph()
    p.text = "for Central Government Portals"
    p.font.size = Pt(18)
    p.font.color.rgb = TEXT_MUTED
    p.space_after = Pt(28)
    
    badge_data = [
        ("Architecture", "4-Tier Microservices (React, Express, FastAPI, Postgres)"),
        ("Machine Learning", "TF-IDF Vectorization + Calibrated Logistic Regression"),
        ("Evaluation Dataset", "675 Academic Synthetic Samples (9 Categories)"),
        ("Verification", "99 Automated Tests Passing (34 Pytest, 65 Node Tests)")
    ]
    
    for idx, (label, val) in enumerate(badge_data):
        bx = 1.2 + (idx % 2) * 5.4
        by = 4.2 + (idx // 2) * 1.1
        card = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(bx), Inches(by), Inches(5.1), Inches(0.9))
        card.fill.solid()
        card.fill.fore_color.rgb = RGBColor(241, 245, 249)
        card.line.color.rgb = CARD_BORDER
        
        ctf = card.text_frame
        ctf.word_wrap = True
        ctf.margin_left = Inches(0.15)
        ctf.margin_top = Inches(0.1)
        cp1 = ctf.paragraphs[0]
        cp1.text = label.upper()
        cp1.font.size = Pt(9)
        cp1.font.bold = True
        cp1.font.color.rgb = BLUE_ACCENT
        
        cp2 = ctf.add_paragraph()
        cp2.text = val
        cp2.font.size = Pt(11)
        cp2.font.bold = True
        cp2.font.color.rgb = TEXT_MAIN

    # ----------------------------------------------------
    # SLIDE 2: Problem Statement (Qualitative Statement - No Unsupported Statistics)
    # ----------------------------------------------------
    s2 = prs.slides.add_slide(blank_layout)
    set_slide_background(s2)
    add_header(s2, "Problem Statement: Central Portal Triage Bottlenecks", "BACKGROUND & MOTIVATION")
    
    add_card(s2, 0.8, 1.5, 5.7, 5.2, "Identified Operational Bottlenecks", NAVY_PRIMARY)
    tb = s2.shapes.add_textbox(Inches(1.0), Inches(2.1), Inches(5.3), Inches(4.4))
    tf = tb.text_frame
    tf.word_wrap = True
    
    probs = [
        ("High Complaint Latency", "Central portals receive thousands of complaints daily; manual triage creates massive processing backlogs."),
        ("Frequent Misrouting", "Citizens manually select target departments from dropdown menus, leading to inter-department transfers."),
        ("Unprioritized Safety Risks", "Urgent infrastructure or public safety hazards wait in line behind routine inquiries."),
        ("Multilingual Barrier", "Citizens file complaints in diverse regional Indic scripts, creating language translation barriers for officers."),
        ("Opaque Citizen Tracking", "Lack of automated contextual status assistance leaves citizens unaware of complaint progress.")
    ]
    for i, (title, desc) in enumerate(probs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = f"• {title}: "
        p.font.bold = True
        p.font.size = Pt(12)
        p.font.color.rgb = BLUE_ACCENT
        
        run = p.add_run()
        run.text = desc
        run.font.bold = False
        run.font.color.rgb = TEXT_MAIN
        run.font.size = Pt(11)
        p.space_after = Pt(8)
        
    add_card(s2, 6.8, 1.5, 5.7, 5.2, "Operational Impact & Target Focus", NAVY_PRIMARY)
    tb2 = s2.shapes.add_textbox(Inches(7.0), Inches(2.1), Inches(5.3), Inches(4.4))
    tf2 = tb2.text_frame
    tf2.word_wrap = True
    
    impacts = [
        ("Triage Latency", "Manual sorting creates processing delays before complaints reach officer queues."),
        ("Routing Accuracy", "High frequency of inter-department misrouting when citizens select target departments manually."),
        ("Emergency Escalation", "Lack of automated keyword detection to escalate urgent leaks, power hazards, or contamination."),
        ("Language Friction", "Non-English regional complaints face severe processing delays due to manual script identification.")
    ]
    for i, (title, desc) in enumerate(impacts):
        p = tf2.paragraphs[0] if i == 0 else tf2.add_paragraph()
        p.text = f"• {title}: "
        p.font.bold = True
        p.font.size = Pt(12)
        p.font.color.rgb = GOLD_AMBER
        
        run = p.add_run()
        run.text = desc
        run.font.bold = False
        run.font.color.rgb = TEXT_MAIN
        run.font.size = Pt(11)
        p.space_after = Pt(12)

    # ----------------------------------------------------
    # SLIDE 3: Objectives
    # ----------------------------------------------------
    s3 = prs.slides.add_slide(blank_layout)
    set_slide_background(s3)
    add_header(s3, "Project Objectives & Technical Scope", "PROJECT GOALS")
    
    objs = [
        ("1. Automated AI Classification", "Implement statistical machine learning (TF-IDF + Logistic Regression) to classify raw complaint text into 9 primary government domains with confidence scores."),
        ("2. Urgency & Priority Assessment", "Analyze emergency keywords to predict complaint priority (LOW, MEDIUM, HIGH, CRITICAL) and assign SLA resolution deadlines."),
        ("3. Confidence-Based Human Review", "Establish automated thresholding (75% confidence) to route uncertain AI predictions to administrators for human classification correction."),
        ("4. Explainable AI (XAI) Transparency", "Extract exact TF-IDF token contribution weights for predictions, providing clear feature explanations for every assigned category."),
        ("5. Role-Based Governance", "Build isolated web portals for Citizens (submit & track), Officers (department-locked queues & status updates), and Admins (analytics & overrides)."),
        ("6. Intent-Based Conversational Chatbot", "Provide a 24/7 multi-turn citizen assistant for status tracking, reference retrieval, filing guidance, and department inquiry.")
    ]
    
    for idx, (title, desc) in enumerate(objs):
        col = idx % 2
        row = idx // 2
        left = 0.8 + col * 5.9
        top = 1.5 + row * 1.8
        
        add_card(s3, left, top, 5.6, 1.6, title, BLUE_ACCENT)
        tb = s3.shapes.add_textbox(Inches(left + 0.2), Inches(top + 0.45), Inches(5.2), Inches(1.0))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = desc
        p.font.size = Pt(11)
        p.font.color.rgb = TEXT_MAIN

    # ----------------------------------------------------
    # SLIDE 4: Existing System vs Proposed System
    # ----------------------------------------------------
    s4 = prs.slides.add_slide(blank_layout)
    set_slide_background(s4)
    add_header(s4, "System Comparison: Existing Manual vs Proposed Framework", "SYSTEM COMPARISON")
    
    table_shape = s4.shapes.add_table(6, 3, Inches(0.8), Inches(1.5), Inches(11.733), Inches(5.2))
    table = table_shape.table
    table.columns[0].width = Inches(2.7)
    table.columns[1].width = Inches(4.5)
    table.columns[2].width = Inches(4.533)
    
    headers = ["Feature / Dimension", "Existing Conventional Portals", "Proposed Intelligent Framework"]
    for i, h in enumerate(headers):
        cell = table.cell(0, i)
        cell.text = h
        cell.fill.solid()
        cell.fill.fore_color.rgb = NAVY_PRIMARY
        for p in cell.text_frame.paragraphs:
            p.font.bold = True
            p.font.size = Pt(13)
            p.font.color.rgb = RGBColor(255, 255, 255)
            
    rows_data = [
        ("Department Triage", "Manual selection by citizen; high error rate.", "Automated TF-IDF + Logistic Regression classification."),
        ("Priority & Urgency", "First-in-first-out (FIFO); no emergency urgency flag.", "Urgency term analysis assigning LOW to CRITICAL SLAs."),
        ("Uncertain Predictions", "Accepted automatically or misrouted silently.", "Confidence threshold check (75%) flagging human review."),
        ("Explainability (XAI)", "Black-box or manual administrative decision.", "Exact TF-IDF token contribution scores returned per class."),
        ("Citizen Support", "Static FAQ pages or slow email queries.", "24/7 multi-turn intent-based chatbot for reference tracking.")
    ]
    
    for r_idx, r_data in enumerate(rows_data, start=1):
        bg = RGBColor(255, 255, 255) if r_idx % 2 == 1 else RGBColor(241, 245, 249)
        for c_idx, val in enumerate(r_data):
            cell = table.cell(r_idx, c_idx)
            cell.text = val
            cell.fill.solid()
            cell.fill.fore_color.rgb = bg
            for p in cell.text_frame.paragraphs:
                p.font.size = Pt(11)
                p.font.color.rgb = TEXT_MAIN
                if c_idx == 0:
                    p.font.bold = True
                elif c_idx == 2:
                    p.font.bold = True
                    p.font.color.rgb = BLUE_ACCENT

    # ----------------------------------------------------
    # SLIDE 5: Overall System Architecture
    # ----------------------------------------------------
    s5 = prs.slides.add_slide(blank_layout)
    set_slide_background(s5)
    add_header(s5, "Decoupled 4-Tier Microservices Architecture", "SYSTEM ARCHITECTURE")
    
    tiers = [
        ("1. React + TS Frontend UI", "Port 80 (Nginx) / 5173", "Single Page App (SPA), AuthContext, Axios API Client, Role-based Dashboards, Tailwind CSS", BLUE_ACCENT),
        ("2. Node.js + Express Gateway", "Port 5000", "JWT Auth, RBAC Middleware, Workflow Logic, Prisma ORM, Notification Triggers", NAVY_PRIMARY),
        ("3. Python FastAPI AI Service", "Port 8001", "FastAPI, Scikit-learn ML Pipeline, Langdetect, Priority Predictor, Chatbot Engine", GREEN_SUCCESS),
        ("4. PostgreSQL 15 Database", "Port 5432 / 5433", "Relational persistence: Users, Departments, Grievances, Status History, Notifications", GOLD_AMBER)
    ]
    
    for i, (title, port, desc, color) in enumerate(tiers):
        top = 1.5 + i * 1.3
        add_card(s5, 0.8, top, 11.733, 1.15, title, color)
        
        pb = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(9.8), Inches(top + 0.15), Inches(2.5), Inches(0.35))
        pb.fill.solid()
        pb.fill.fore_color.rgb = RGBColor(241, 245, 249)
        pb.line.color.rgb = CARD_BORDER
        ptf = pb.text_frame
        pp = ptf.paragraphs[0]
        pp.text = port
        pp.font.size = Pt(10)
        pp.font.bold = True
        pp.font.color.rgb = TEXT_MUTED
        
        tb = s5.shapes.add_textbox(Inches(1.0), Inches(top + 0.45), Inches(8.5), Inches(0.6))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = desc
        p.font.size = Pt(11)
        p.font.color.rgb = TEXT_MAIN

    # ----------------------------------------------------
    # SLIDE 6: Technology Stack
    # ----------------------------------------------------
    s6 = prs.slides.add_slide(blank_layout)
    set_slide_background(s6)
    add_header(s6, "Technology Stack & Core Dependencies", "TECH STACK")
    
    techs = [
        ("Frontend UI Layer", [("Framework", "React 18 + TypeScript"), ("Build Tool", "Vite"), ("Styling", "Tailwind CSS"), ("Icons", "Lucide React")]),
        ("Backend Gateway Layer", [("Runtime", "Node.js v20+"), ("Framework", "Express.js v4"), ("ORM", "Prisma ORM v6"), ("Security", "JWT + bcryptjs")]),
        ("AI Microservice Layer", [("Language", "Python 3.11 / 3.13"), ("Framework", "FastAPI"), ("ML Toolkit", "Scikit-learn + Joblib"), ("NLP", "Langdetect")]),
        ("Database & DevOps", [("Database", "PostgreSQL 15"), ("Containerization", "Docker & Docker Compose"), ("Testing", "Pytest & Node --test"), ("Quality", "ESLint & Vite Build")])
    ]
    
    for idx, (category, items) in enumerate(techs):
        col = idx % 2
        row = idx // 2
        left = 0.8 + col * 5.9
        top = 1.5 + row * 2.7
        
        add_card(s6, left, top, 5.6, 2.5, category, BLUE_ACCENT)
        tb = s6.shapes.add_textbox(Inches(left + 0.2), Inches(top + 0.5), Inches(5.2), Inches(1.9))
        tf = tb.text_frame
        tf.word_wrap = True
        
        for i, (k, v) in enumerate(items):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.text = f"• {k}: "
            p.font.bold = True
            p.font.size = Pt(11)
            p.font.color.rgb = NAVY_PRIMARY
            
            run = p.add_run()
            run.text = v
            run.font.bold = False
            run.font.size = Pt(11)
            run.font.color.rgb = TEXT_MAIN
            p.space_after = Pt(4)

    # ----------------------------------------------------
    # SLIDE 7: Database Design & Schema Models
    # ----------------------------------------------------
    s7 = prs.slides.add_slide(blank_layout)
    set_slide_background(s7)
    add_header(s7, "Database Design: Prisma Relational Schema Models", "DATABASE SCHEMA")
    
    models = [
        ("User", "Stores account credentials, role (CITIZEN, OFFICER, ADMIN), and optional department_id for officers."),
        ("Department", "Master catalog of 9 government departments (Water Supply, Electricity, PWD, Healthcare, etc.)."),
        ("Grievance", "Central entity storing complaint text, status, priority, department, AI confidence, XAI terms, and human correction flags."),
        ("GrievanceStatusHistory", "Immutable audit trail capturing status transitions (SUBMITTED → UNDER_REVIEW → RESOLVED/REJECTED) and officer remarks."),
        ("Notification", "In-app citizen notification records triggered on workflow events (STATUS_CHANGED, REVIEW_FLAGGED).")
    ]
    
    for i, (name, desc) in enumerate(models):
        top = 1.5 + i * 1.05
        add_card(s7, 0.8, top, 11.733, 0.95, f"Model: {name}", NAVY_PRIMARY)
        tb = s7.shapes.add_textbox(Inches(2.8), Inches(top + 0.15), Inches(9.5), Inches(0.65))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = desc
        p.font.size = Pt(11)
        p.font.color.rgb = TEXT_MAIN

    # ----------------------------------------------------
    # SLIDE 8: End-to-End Grievance Workflow
    # ----------------------------------------------------
    s8 = prs.slides.add_slide(blank_layout)
    set_slide_background(s8)
    add_header(s8, "End-to-End Grievance Processing & Resolution Workflow", "WORKFLOW LIFECYCLE")
    
    steps = [
        ("1. Submission", "Citizen enters complaint text (min 5 chars)."),
        ("2. Language & Prep", "Langdetect checks script; text lowercased."),
        ("3. ML Inference", "TF-IDF + LogReg predicts category & confidence."),
        ("4. Urgency Analysis", "Urgency keywords set priority & SLA hours."),
        ("5. Confidence Check", "Conf < 75% flags ai_review_required: true."),
        ("6. DB Persistence", "Saved in Postgres; reference GRV-2026-XXXX issued."),
        ("7. Officer Review", "Department officer updates status + remarks."),
        ("8. Citizen Notice", "In-app status change notification sent.")
    ]
    
    for idx, (title, desc) in enumerate(steps):
        col = idx % 4
        row = idx // 4
        left = 0.8 + col * 2.95
        top = 1.6 + row * 2.6
        
        add_card(s8, left, top, 2.75, 2.3, title, BLUE_ACCENT)
        tb = s8.shapes.add_textbox(Inches(left + 0.15), Inches(top + 0.5), Inches(2.45), Inches(1.6))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = desc
        p.font.size = Pt(10)
        p.font.color.rgb = TEXT_MAIN

    # ----------------------------------------------------
    # SLIDE 9: AI/ML Methodology
    # ----------------------------------------------------
    s9 = prs.slides.add_slide(blank_layout)
    set_slide_background(s9)
    add_header(s9, "AI/ML Methodology: TF-IDF + Logistic Regression", "MACHINE LEARNING")
    
    add_card(s9, 0.8, 1.5, 5.7, 5.2, "TF-IDF Feature Representation", BLUE_ACCENT)
    tb = s9.shapes.add_textbox(Inches(1.0), Inches(2.1), Inches(5.3), Inches(4.4))
    tf = tb.text_frame
    tf.word_wrap = True
    
    v_specs = [
        ("N-gram Range", "ngram_range=(1, 2) (Unigrams & Bigrams)"),
        ("Max Features", "max_features=5000 top informative terms"),
        ("Sublinear TF", "sublinear_tf=True (1 + log(tf) scaling)"),
        ("Preprocessing", "lowercase=True, stop-word & noise removal"),
        ("Mathematical TF-IDF", "TF-IDF(t,d) = TF(t,d) x log(N / df_t)")
    ]
    for i, (k, v) in enumerate(v_specs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = f"• {k}: "
        p.font.bold = True
        p.font.size = Pt(11)
        p.font.color.rgb = NAVY_PRIMARY
        run = p.add_run()
        run.text = v
        run.font.bold = False
        run.font.size = Pt(11)
        run.font.color.rgb = TEXT_MAIN
        p.space_after = Pt(8)

    add_card(s9, 6.8, 1.5, 5.7, 5.2, "Logistic Regression & XAI Explanations", NAVY_PRIMARY)
    tb2 = s9.shapes.add_textbox(Inches(7.0), Inches(2.1), Inches(5.3), Inches(4.4))
    tf2 = tb2.text_frame
    tf2.word_wrap = True
    
    c_specs = [
        ("Inverse Reg Strength", "C=2.5 (optimal fit / regularization balance)"),
        ("Solver & Iterations", "solver='lbfgs', max_iter=1000, seed=42"),
        ("Softmax Multiclass", "P(y=k|x) = exp(w_k.x + b_k) / Sum(exp(w_j.x + b_j))"),
        ("XAI Token Weights", "Score(t) = TF-IDF(t) x w_predicted_class(t)"),
        ("Academic Disclaimer", "Does NOT use Transformer models or LLMs")
    ]
    for i, (k, v) in enumerate(c_specs):
        p = tf2.paragraphs[0] if i == 0 else tf2.add_paragraph()
        p.text = f"• {k}: "
        p.font.bold = True
        p.font.size = Pt(11)
        p.font.color.rgb = BLUE_ACCENT
        run = p.add_run()
        run.text = v
        run.font.bold = False
        run.font.size = Pt(11)
        run.font.color.rgb = TEXT_MAIN
        p.space_after = Pt(8)

    # ----------------------------------------------------
    # SLIDE 10: Confidence & Human-in-the-Loop Review
    # ----------------------------------------------------
    s10 = prs.slides.add_slide(blank_layout)
    set_slide_background(s10)
    add_header(s10, "Confidence Thresholding & Human Classification Review", "HUMAN-IN-THE-LOOP")
    
    threshs = [
        ("HIGH Confidence (>= 0.75)", "Automatic Dispatch", "ai_review_required: false. Complaint is routed directly to the predicted department queue.", GREEN_SUCCESS),
        ("MEDIUM Confidence (0.50 - 0.74)", "Flagged for Human Review", "ai_review_required: true. Complaint is flagged in Admin dashboard for manual inspection.", GOLD_AMBER),
        ("LOW Confidence (< 0.50)", "Escalated Human Review", "ai_review_required: true. Low-confidence or out-of-domain text assigned category 'Other'.", RGBColor(185, 28, 28))
    ]
    
    for i, (title, tag, desc, color) in enumerate(threshs):
        top = 1.5 + i * 1.7
        add_card(s10, 0.8, top, 11.733, 1.5, title, color)
        
        tb_shape = s10.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(9.2), Inches(top + 0.15), Inches(3.1), Inches(0.35))
        tb_shape.fill.solid()
        tb_shape.fill.fore_color.rgb = RGBColor(241, 245, 249)
        tb_shape.line.color.rgb = CARD_BORDER
        ptf = tb_shape.text_frame
        pp = ptf.paragraphs[0]
        pp.text = tag.upper()
        pp.font.size = Pt(10)
        pp.font.bold = True
        pp.font.color.rgb = color
        
        tb = s10.shapes.add_textbox(Inches(1.0), Inches(top + 0.5), Inches(11.3), Inches(0.8))
        tf = tb.text_frame
        tf.word_wrap = True
        
        p = tf.paragraphs[0]
        p.text = desc
        p.font.size = Pt(11)
        p.font.color.rgb = TEXT_MAIN
        p.space_after = Pt(4)
        
        p2 = tf.add_paragraph()
        p2.text = "Audit Trail: Admin correction updates category & department_id, preserving ai_original_category and setting is_human_corrected: true."
        p2.font.size = Pt(10)
        p2.font.bold = True
        p2.font.color.rgb = TEXT_MUTED

    # ----------------------------------------------------
    # SLIDE 11: Multilingual Support & Chatbot
    # ----------------------------------------------------
    s11 = prs.slides.add_slide(blank_layout)
    set_slide_background(s11)
    add_header(s11, "Multilingual Language Detection & Intent Chatbot", "MULTILINGUAL & CHATBOT")
    
    add_card(s11, 0.8, 1.5, 5.7, 5.2, "Multilingual Handling (Langdetect)", BLUE_ACCENT)
    tb = s11.shapes.add_textbox(Inches(1.0), Inches(2.1), Inches(5.3), Inches(4.4))
    tf = tb.text_frame
    tf.word_wrap = True
    
    m_facts = [
        ("Language Detector", "Uses langdetect library to identify ISO-639-1 language codes (en, hi, ta, te, mr, etc.)."),
        ("Script Preservation", "Preserves original regional text in Postgres without altering citizen input."),
        ("Keyword Fallback", "Applies keyword rules for regional script domain terms."),
        ("Honest Disclaimer", "Does NOT perform IndicTrans2 neural translation; regional text preserved as-is.")
    ]
    for i, (k, v) in enumerate(m_facts):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = f"• {k}: "
        p.font.bold = True
        p.font.size = Pt(11)
        p.font.color.rgb = NAVY_PRIMARY
        run = p.add_run()
        run.text = v
        run.font.bold = False
        run.font.size = Pt(11)
        run.font.color.rgb = TEXT_MAIN
        p.space_after = Pt(8)

    add_card(s11, 6.8, 1.5, 5.7, 5.2, "Intent-Based Citizen Chatbot (9 Intents)", NAVY_PRIMARY)
    tb2 = s11.shapes.add_textbox(Inches(7.0), Inches(2.1), Inches(5.3), Inches(4.4))
    tf2 = tb2.text_frame
    tf2.word_wrap = True
    
    c_facts = [
        ("9 Explicit Intents", "GREETING, PRIORITY_INFO, PROCESS_INFO, TRACKING_GUIDANCE, SUBMIT_GUIDANCE, DEPT_INFO, HELP, GRIEVANCE_STATUS, UNKNOWN."),
        ("Entity Extraction", "Regex matches GRV-YYYY-XXXXXX reference tracking codes."),
        ("Context Recall", "Session context memory recalls last_grievance_number for pronouns ('it', 'this complaint')."),
        ("Data Privacy", "Filters queries by authenticated user's grievances (user_id == req.user.id).")
    ]
    for i, (k, v) in enumerate(c_facts):
        p = tf2.paragraphs[0] if i == 0 else tf2.add_paragraph()
        p.text = f"• {k}: "
        p.font.bold = True
        p.font.size = Pt(11)
        p.font.color.rgb = BLUE_ACCENT
        run = p.add_run()
        run.text = v
        run.font.bold = False
        run.font.size = Pt(11)
        run.font.color.rgb = TEXT_MAIN
        p.space_after = Pt(8)

    # ----------------------------------------------------
    # SLIDE 12: Security & Access Control
    # ----------------------------------------------------
    s12 = prs.slides.add_slide(blank_layout)
    set_slide_background(s12)
    add_header(s12, "Security Architecture & Role-Based Access Control", "SECURITY & PRIVACY")
    
    secs = [
        ("JWT Authentication", "Stateless 24-hour JSON Web Tokens passed in Authorization: Bearer header. Validated via authenticate middleware."),
        ("Password Security", "Salted bcrypt hashing (cost factor 10) applied to passwords prior to PostgreSQL persistence."),
        ("Role-Based Access Control", "requireRole('CITIZEN', 'OFFICER', 'ADMIN') middleware blocks unauthorized endpoint invocation with HTTP 403."),
        ("Citizen Ownership Isolation", "Grievance lookups strictly include where: { user_id: req.user.id }, preventing cross-user data access."),
        ("Officer Department Isolation", "Officer queues filter by where: { department_id: req.user.department_id }, enforcing department boundaries."),
        ("SQL Injection Protection", "Prisma ORM parameterizes all SQL database queries, eliminating raw SQL string injection risks.")
    ]
    
    for idx, (title, desc) in enumerate(secs):
        col = idx % 2
        row = idx // 2
        left = 0.8 + col * 5.9
        top = 1.5 + row * 1.8
        
        add_card(s12, left, top, 5.6, 1.6, title, BLUE_ACCENT)
        tb = s12.shapes.add_textbox(Inches(left + 0.2), Inches(top + 0.45), Inches(5.2), Inches(1.0))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = desc
        p.font.size = Pt(11)
        p.font.color.rgb = TEXT_MAIN

    # ----------------------------------------------------
    # SLIDE 13: Application Screenshots Showcase (Embedded Genuine PNG Images)
    # ----------------------------------------------------
    s13 = prs.slides.add_slide(blank_layout)
    set_slide_background(s13)
    add_header(s13, "Application User Interface & Portal Demonstration", "UI SHOWCASE")
    
    screens = [
        ("1. Citizen Submission", "01_citizen_submission.png"),
        ("2. Multilingual Chatbot", "02_chatbot_interaction.png"),
        ("3. Officer Queue", "03_officer_queue.png"),
        ("4. Admin Overview", "04_admin_dashboard.png")
    ]
    
    for idx, (title, img_filename) in enumerate(screens):
        col = idx % 2
        row = idx // 2
        left = 0.8 + col * 5.9
        top = 1.4 + row * 2.8
        
        # Card container
        add_card(s13, left, top, 5.6, 2.65, title, NAVY_PRIMARY)
        
        img_path = os.path.join("docs", "presentation", "screenshots", img_filename)
        if os.path.exists(img_path):
            s13.shapes.add_picture(img_path, Inches(left + 0.15), Inches(top + 0.5), width=Inches(5.3), height=Inches(2.0))

    # ----------------------------------------------------
    # SLIDE 14: Testing & Experimental Results
    # ----------------------------------------------------
    s14 = prs.slides.add_slide(blank_layout)
    set_slide_background(s14)
    add_header(s14, "Empirical Testing & Verification Metrics", "TESTING RESULTS")
    
    metrics = [
        ("34 / 34 Passed", "Python Pytest AI Suite", GREEN_SUCCESS),
        ("65 / 65 Passed", "Node Backend Suite", GREEN_SUCCESS),
        ("0 Errors", "Frontend ESLint Quality", BLUE_ACCENT),
        ("SUCCESS", "Vite Production Build", BLUE_ACCENT)
    ]
    
    for i, (val, label, col) in enumerate(metrics):
        left = 0.8 + i * 2.95
        add_card(s14, left, 1.5, 2.75, 1.3, None, bg_color=CARD_BG)
        
        tb = s14.shapes.add_textbox(Inches(left + 0.1), Inches(1.6), Inches(2.55), Inches(1.1))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = val
        p.alignment = PP_ALIGN.CENTER
        p.font.size = Pt(18)
        p.font.bold = True
        p.font.color.rgb = col
        
        p2 = tf.add_paragraph()
        p2.text = label
        p2.alignment = PP_ALIGN.CENTER
        p2.font.size = Pt(10)
        p2.font.color.rgb = TEXT_MAIN

    add_card(s14, 0.8, 3.1, 11.733, 3.6, "Empirical ML Evaluation: Clean Benchmark vs Realistic Edge Cases", NAVY_PRIMARY)
    tb = s14.shapes.add_textbox(Inches(1.0), Inches(3.6), Inches(11.333), Inches(2.9))
    tf = tb.text_frame
    tf.word_wrap = True
    
    eval_bullets = [
        ("Clean Benchmark Split", "675 synthetic samples split strictly into Train (472, 70%), Validation (101, 15%), and Test (102, 15%). Seed: 42."),
        ("Clean Test Accuracy", "100.0% accuracy, precision, recall, and F1-score on clean held-out test split (102 samples) via evaluate_classifier.py."),
        ("Realistic Edge-Case Test", "79.63% accuracy (86/108 correct) on 108 realistic validation cases with typos, slang, and multi-issue text via error_analysis.py."),
        ("Human Review Safety Net", "88.9% (96/108) of ambiguous validation cases safely flagged for Human Review (confidence < 0.75), preventing automated misrouting."),
        ("Academic Honesty", "Synthetic templates achieve 100%, but noisy inputs reduce direct ML accuracy to ~80%, proving the necessity of confidence thresholding.")
    ]
    for i, (k, v) in enumerate(eval_bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = f"• {k}: "
        p.font.bold = True
        p.font.size = Pt(10.5)
        p.font.color.rgb = BLUE_ACCENT
        run = p.add_run()
        run.text = v
        run.font.bold = False
        run.font.size = Pt(10.5)
        run.font.color.rgb = TEXT_MAIN
        p.space_after = Pt(4)

    # ----------------------------------------------------
    # SLIDE 15: Conclusion, Limitations & Future Work
    # ----------------------------------------------------
    s15 = prs.slides.add_slide(blank_layout)
    set_slide_background(s15)
    add_header(s15, "Conclusion, Verified Limitations & Future Roadmap", "CONCLUSION & FUTURE SCOPE")
    
    add_card(s15, 0.8, 1.5, 5.7, 5.2, "Verified Technical Limitations", RGBColor(185, 28, 28))
    tb = s15.shapes.add_textbox(Inches(1.0), Inches(2.1), Inches(5.3), Inches(4.4))
    tf = tb.text_frame
    tf.word_wrap = True
    
    lims = [
        ("Classifier Scope", "Uses TF-IDF + Logistic Regression (does NOT use BERT or deep Transformers)."),
        ("Language Limits", "Language detection active; regional text preserved without neural IndicTrans2 translation."),
        ("Chatbot Scope", "Explicit intent matching schemas; not a generative LLM or RAG model."),
        ("Synthetic Dataset", "Evaluated on 675 synthetic samples; real-world government data required for production scale.")
    ]
    for i, (k, v) in enumerate(lims):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = f"• {k}: "
        p.font.bold = True
        p.font.size = Pt(11)
        p.font.color.rgb = NAVY_PRIMARY
        run = p.add_run()
        run.text = v
        run.font.bold = False
        run.font.size = Pt(11)
        run.font.color.rgb = TEXT_MAIN
        p.space_after = Pt(8)

    add_card(s15, 6.8, 1.5, 5.7, 5.2, "Future Technical Enhancements", GREEN_SUCCESS)
    tb2 = s15.shapes.add_textbox(Inches(7.0), Inches(2.1), Inches(5.3), Inches(4.4))
    tf2 = tb2.text_frame
    tf2.word_wrap = True
    
    futures = [
        ("IndicBERT Transformer", "Fine-tune IndicBERT or mBERT embeddings for contextual Indic language understanding."),
        ("Neural Translation", "Integrate IndicTrans2 NMT model for 22 official Indian languages."),
        ("Active Learning Pipeline", "Collect admin human corrections to automatically retrain model weights."),
        ("Telephony & SMS", "Connect Twilio / NIC SMS & WhatsApp gateways for low-literacy populations.")
    ]
    for i, (k, v) in enumerate(futures):
        p = tf2.paragraphs[0] if i == 0 else tf2.add_paragraph()
        p.text = f"• {k}: "
        p.font.bold = True
        p.font.size = Pt(11)
        p.font.color.rgb = BLUE_ACCENT
        run = p.add_run()
        run.text = v
        run.font.bold = False
        run.font.size = Pt(11)
        run.font.color.rgb = TEXT_MAIN
        p.space_after = Pt(8)

    output_dir = os.path.join("docs", "presentation")
    os.makedirs(output_dir, exist_ok=True)
    output_path = os.path.join(output_dir, "final-project-presentation.pptx")
    
    prs.save(output_path)
    print(f"Successfully generated PowerPoint presentation at: {output_path}")

if __name__ == "__main__":
    create_presentation()
